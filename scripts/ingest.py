#!/usr/bin/env python3
"""Extract locally supplied official documents into a review queue, never public data.

This tool does not retrieve documents, infer their economic scope, or approve values.
Every candidate retains its document checksum and an exact decimal string.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
import sys
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

SOURCE_DOMAINS = {
    "mef": "finances.gov.ma", "tgr": "tgr.gov.ma", "bam": "bkam.ma",
    "hcp": "hcp.ma", "oc": "oc.gov.ma", "cour": "courdescomptes.ma",
    "opendata": "data.gov.ma", "imf": "imf.org", "worldbank": "worldbank.org",
}
INDICATORS = {
    "debt_treasury": ("MAD", "treasury"),
    "debt_general_government": ("MAD", "general_government"),
    "debt_domestic": ("MAD", "treasury"), "debt_external": ("MAD", "treasury"),
    "debt_gdp": ("percent", "treasury"), "debt_per_capita": ("MAD", "treasury"),
    "debt_stock_change": ("MAD", "treasury"), "interest": ("MAD", "treasury"),
    "deficit": ("MAD", "central_government"), "revenue": ("MAD", "central_government"),
    "expenditure": ("MAD", "central_government"), "debt_service": ("MAD", "treasury"),
    "population": ("persons", "resident_population"), "gdp": ("MAD", "national_economy"),
}
MAX_INPUT_BYTES = 32 * 1024 * 1024
MAX_ROWS = 50_000
REQUIRED_COLUMNS = {"indicatorId", "exactValue", "unit", "scale", "scope", "period", "referenceDate", "status"}


class IngestionError(ValueError):
    """A source file, manifest or candidate failed review-queue validation."""


def require(value: Any, name: str) -> str:
    if value is None or not str(value).strip():
        raise IngestionError(f"Required field: {name}")
    return str(value).strip()


def calendar_date(value: Any, name: str) -> str:
    text = require(value, name)
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        raise IngestionError(f"{name} must be an ISO calendar date")
    try:
        date.fromisoformat(text)
    except ValueError as error:
        raise IngestionError(f"Invalid {name}: {text}") from error
    return text


def normalize_decimal(value: Any) -> str:
    """Accept decimal dot or French decimal comma; never guess comma grouping."""
    text = require(value, "exactValue")
    text = re.sub(r"[\s\u00a0\u202f]", "", text)
    if "," in text and "." in text:
        raise IngestionError("Ambiguous separators: explicitly normalize the source value")
    text = text.replace(",", ".").replace("−", "-")
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", text):
        raise IngestionError(f"Invalid decimal: {value!r}")
    try:
        number = Decimal(text)
    except InvalidOperation as error:
        raise IngestionError("Invalid decimal") from error
    if not number.is_finite() or not math.isfinite(float(number)):
        raise IngestionError("Numeric value exceeds supported finite range")
    return format(number, "f")


def source_document(manifest: dict[str, Any], input_path: Path, contents: bytes) -> dict[str, Any]:
    document = manifest.get("document")
    if not isinstance(document, dict):
        raise IngestionError("The manifest must contain a document object")
    source_id = require(document.get("sourceId"), "document.sourceId")
    if source_id not in SOURCE_DOMAINS:
        raise IngestionError("Unknown source institution")
    url = require(document.get("sourceUrl"), "document.sourceUrl")
    parsed = urlparse(url)
    domain = SOURCE_DOMAINS[source_id]
    if parsed.scheme != "https" or parsed.username or parsed.password or not (
        parsed.hostname == domain or (parsed.hostname or "").endswith("." + domain)
    ):
        raise IngestionError("Source URL must use HTTPS on the institution’s official domain")
    digest = hashlib.sha256(contents).hexdigest()
    expected = document.get("expectedSha256")
    if expected is not None and expected != digest:
        raise IngestionError("Document SHA-256 does not match expectedSha256")
    title = require(document.get("title"), "document.title")
    if len(title) < 3:
        raise IngestionError("Document title must have at least three characters")
    return {
        "sourceId": source_id, "sourceUrl": url,
        "title": title,
        "publicationDate": calendar_date(document.get("publicationDate"), "publicationDate"),
        "filename": input_path.name, "sha256": digest,
        "byteSize": len(contents), "authenticityVerified": False,
    }


def candidate(row: dict[str, Any], document: dict[str, Any], locator: str, page: int | None = None,
              excerpt: str | None = None) -> dict[str, Any]:
    missing = REQUIRED_COLUMNS - row.keys()
    if missing:
        raise IngestionError("Missing columns: " + ", ".join(sorted(missing)))
    indicator = require(row.get("indicatorId"), "indicatorId")
    if indicator not in INDICATORS:
        raise IngestionError(f"Unknown indicator: {indicator}")
    unit, scope = INDICATORS[indicator]
    if row.get("unit") != unit or row.get("scope") != scope:
        raise IngestionError(f"{indicator} requires unit={unit} and scope={scope}; no conversion is inferred")
    try:
        scale = int(require(row.get("scale"), "scale"))
    except ValueError as error:
        raise IngestionError("Scale must be an integer") from error
    if scale not in (1, 1_000_000, 1_000_000_000) or (unit == "percent" and scale != 1):
        raise IngestionError("Unsupported scale")
    exact = normalize_decimal(row.get("exactValue"))
    number = Decimal(exact)
    if indicator not in ("deficit", "debt_stock_change") and number < 0:
        raise IngestionError(f"{indicator} cannot be negative")
    if indicator in ("population", "gdp") and number <= 0:
        raise IngestionError(f"{indicator} must be positive")
    status = require(row.get("status"), "status")
    if status not in ("observed", "provisional"):
        raise IngestionError("Only observed or provisional source data belong in the observation review queue")
    reference = calendar_date(row.get("referenceDate"), "referenceDate")
    if reference > document["publicationDate"]:
        raise IngestionError("Observed data cannot be published before its reference date")
    period = require(row.get("period"), "period")
    raw_page = row.get("documentPage")
    if raw_page not in (None, ""):
        try:
            page = int(raw_page)
        except (ValueError, TypeError) as error:
            raise IngestionError("documentPage must be a positive integer") from error
    if page is not None and page <= 0:
        raise IngestionError("documentPage must be a positive integer")
    identity = "|".join((document["sha256"], indicator, scope, reference, period, locator))
    return {
        "id": "candidate_" + hashlib.sha256(identity.encode()).hexdigest()[:24],
        "indicatorId": indicator, "exactValue": exact, "value": float(number),
        "unit": unit, "scale": scale, "scope": scope, "period": period,
        "referenceDate": reference, "publicationDate": document["publicationDate"],
        "sourceId": document["sourceId"], "sourceUrl": document["sourceUrl"],
        "documentTitle": document["title"], "documentPage": page,
        "documentLocator": str(row.get("documentLocator") or locator),
        "status": status, "validationStatus": "pending_review", "version": 1,
        "provenance": {
            "documentSha256": document["sha256"], "originalValue": str(row.get("exactValue")),
            "extractionLocator": locator, "excerpt": excerpt,
        },
    }


def table_candidates(path: Path, contents: bytes, manifest: dict[str, Any], document: dict[str, Any]) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    rows: list[tuple[dict[str, Any], str]] = []
    if suffix == ".csv":
        try:
            decoded = contents.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise IngestionError("CSV must use UTF-8 encoding") from error
        delimiter = manifest.get("delimiter", ",")
        if not isinstance(delimiter, str) or len(delimiter) != 1:
            raise IngestionError("CSV delimiter must be one character")
        reader = csv.DictReader(io.StringIO(decoded), delimiter=delimiter)
        if reader.fieldnames is None or len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise IngestionError("CSV headers must exist and be unique")
        for number, row in enumerate(reader, 2):
            if None in row:
                raise IngestionError(f"Extra CSV fields at row {number}")
            if number > MAX_ROWS + 1:
                raise IngestionError("Row limit exceeded")
            rows.append((row, f"CSV row {number}"))
    elif suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as error:
            raise IngestionError("Install the pinned Python ingestion requirements for XLSX") from error
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise IngestionError("Uncompressed spreadsheet exceeds 100 MiB")
        book = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=False)
        try:
            sheet_name = require(manifest.get("sheet"), "sheet")
            if sheet_name not in book.sheetnames:
                raise IngestionError("Configured XLSX sheet is missing")
            sheet = book[sheet_name]
            iterator = sheet.iter_rows()
            first = next(iterator, None)
            if first is None:
                raise IngestionError("Empty spreadsheet")
            headers = [str(cell.value) if cell.value is not None else "" for cell in first]
            if "" in headers or len(headers) != len(set(headers)):
                raise IngestionError("Spreadsheet headers must exist and be unique")
            for number, cells in enumerate(iterator, 2):
                if number > MAX_ROWS + 1:
                    raise IngestionError("Row limit exceeded")
                if all(cell.value is None for cell in cells):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise IngestionError(f"Formula cells are not accepted (row {number}); supply the published values")
                row = {header: cell.value for header, cell in zip(headers, cells)}
                rows.append((row, f"XLSX sheet {sheet_name}, row {number}"))
        finally:
            book.close()
    else:
        raise IngestionError("Table input must be .csv or .xlsx")
    return [candidate(row, document, locator) for row, locator in rows]


def pdf_candidates(path: Path, contents: bytes, manifest: dict[str, Any], document: dict[str, Any]) -> list[dict[str, Any]]:
    if path.suffix.lower() != ".pdf":
        raise IngestionError("PDF input must have a .pdf extension")
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise IngestionError("Install the pinned Python ingestion requirements for PDF") from error
    entries = manifest.get("entries")
    if not isinstance(entries, list) or not entries:
        raise IngestionError("PDF mapping requires an explicit entries array")
    if len(entries) > MAX_ROWS:
        raise IngestionError("PDF entry limit exceeded")
    reader = PdfReader(io.BytesIO(contents))
    if reader.is_encrypted:
        raise IngestionError("Encrypted PDF is not supported")
    results = []
    cache: dict[int, str] = {}
    for entry in entries:
        if not isinstance(entry, dict):
            raise IngestionError("Each PDF entry must be an object")
        page = entry.get("documentPage")
        if not isinstance(page, int) or isinstance(page, bool) or not 1 <= page <= len(reader.pages):
            raise IngestionError("Each PDF entry requires a valid one-based documentPage")
        text = cache.setdefault(page, reader.pages[page - 1].extract_text() or "")
        if not text.strip():
            raise IngestionError(f"No text on PDF page {page}; OCR and manual verification are required")
        try:
            pattern = re.compile(require(entry.get("pattern"), "pattern"))
        except re.error as error:
            raise IngestionError("Invalid PDF extraction regex") from error
        if "value" not in pattern.groupindex:
            raise IngestionError("PDF regex must contain the named capture (?P<value>...)")
        matches = list(pattern.finditer(text))
        if len(matches) != 1:
            raise IngestionError(f"PDF page {page}: expected one unambiguous match, found {len(matches)}")
        match = matches[0]
        row = {**entry, "exactValue": match.group("value")}
        excerpt = text[max(0, match.start() - 120):match.end() + 120]
        results.append(candidate(row, document, f"PDF page {page}, match {match.start()}:{match.end()}", page, excerpt))
    return results


def extract(path: Path, manifest: dict[str, Any], kind: str) -> dict[str, Any]:
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise IngestionError("Input exceeds 32 MiB")
    contents = path.read_bytes()
    document = source_document(manifest, path, contents)
    candidates = pdf_candidates(path, contents, manifest, document) if kind == "pdf" else table_candidates(path, contents, manifest, document)
    if not candidates:
        raise IngestionError("No observations were extracted")
    keys: set[tuple[str, ...]] = set()
    for row in candidates:
        key = (row["indicatorId"], row["scope"], row["referenceDate"], row["period"])
        if key in keys:
            raise IngestionError("Duplicate observation in a single document import")
        keys.add(key)
    return {
        "schemaVersion": 1,
        "createdAt": datetime.now(timezone.utc).isoformat(),
        "status": "pending_review",
        "document": document, "candidateCount": len(candidates), "candidates": candidates,
        "reviewRequired": [
            "Verify the document origin and publication date against the official publication.",
            "Check every value, unit, scale, economic scope and reference period against its precise locator.",
            "Determine the document’s revisions, sign conventions and comparability before approval.",
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("kind", choices=("table", "pdf"))
    parser.add_argument("input", type=Path)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    arguments = parser.parse_args()
    try:
        if arguments.output.exists():
            raise IngestionError("Output already exists; choose a new job filename")
        if arguments.output.resolve() == arguments.input.resolve():
            raise IngestionError("Output cannot overwrite the source")
        manifest = json.loads(arguments.manifest.read_text(encoding="utf-8"))
        if not isinstance(manifest, dict):
            raise IngestionError("Manifest must be a JSON object")
        job = extract(arguments.input, manifest, arguments.kind)
        arguments.output.parent.mkdir(parents=True, exist_ok=True)
        with arguments.output.open("x", encoding="utf-8") as output:
            json.dump(job, output, ensure_ascii=False, indent=2, allow_nan=False)
            output.write("\n")
        print(f"{job['candidateCount']} candidate(s), pending review. Saved: {arguments.output}")
        return 0
    except (OSError, ValueError, csv.Error, zipfile.BadZipFile) as error:
        print(f"Ingestion rejected: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
